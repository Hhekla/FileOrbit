#!/usr/bin/env python3
"""Convert actual non-private files through the shipping CLI and inspect their contents.

This intentionally does not count successful process exits as successful conversions.
All fixture files, CLI responses, rendered PDF pages and failed outputs are retained.
Requires the Codex bundled Python (python-docx, pypdf, Pillow), macOS textutil,
and pdftoppm on PATH. It never deletes files and never uploads fixture contents.
"""
import argparse
import hashlib
import html
import json
import re
import shutil
import subprocess
import time
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image
from pypdf import PdfReader


class TextHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.skip = 0
    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style', 'head'):
            self.skip += 1
        elif tag in ('p', 'div', 'li', 'br', 'h1', 'h2'):
            self.parts.append('\n')
    def handle_endtag(self, tag):
        if tag in ('script', 'style', 'head'):
            self.skip = max(0, self.skip - 1)
        elif tag in ('p', 'div', 'li', 'h1', 'h2'):
            self.parts.append('\n')
    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def normalize(text):
    # macOS PDF ToUnicode tables can use compatibility radicals such as ⽂ for 文.
    # Record exact matching separately; do not confuse this with lost paragraphs.
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', text))


def run(args, timeout=30):
    p = subprocess.run(list(map(str, args)), capture_output=True, timeout=timeout)
    return {'args': list(map(str, args)), 'returncode': p.returncode,
            'stdout': p.stdout.decode('utf8', errors='replace'),
            'stderr': p.stderr.decode('utf8', errors='replace')}


def text_of(path, fmt):
    data = path.read_bytes()
    details = {'size': len(data)}
    if fmt == 'pdf':
        assert data.startswith(b'%PDF-'), 'Missing PDF signature'
        pdf = PdfReader(path)
        details['page_count'] = len(pdf.pages)
        assert pdf.pages, 'Zero PDF pages'
        return '\n'.join(p.extract_text() or '' for p in pdf.pages), details
    if fmt in ('docx', 'odt'):
        assert data.startswith(b'PK'), 'Missing ZIP signature'
        member = 'word/document.xml' if fmt == 'docx' else 'content.xml'
        with zipfile.ZipFile(path) as z:
            assert z.testzip() is None, 'Invalid ZIP CRC'
            root = ET.fromstring(z.read(member))
            if fmt == 'docx':
                ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                paragraphs = [''.join(p.itertext()) for p in root.findall('.//w:p', ns)]
                details['bold_runs'] = len(root.findall('.//w:rPr/w:b', ns))
                details['paragraph_count'] = len(paragraphs)
                return '\n'.join(paragraphs), details
            ns = {'t': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'}
            paragraphs = [''.join(p.itertext()) for p in root.iter()
                          if p.tag in ('{' + ns['t'] + '}p', '{' + ns['t'] + '}h')]
            details['paragraph_count'] = len(paragraphs)
            return '\n'.join(paragraphs), details
    if fmt in ('rtf', 'doc'):
        if fmt == 'rtf':
            assert data.startswith(b'{\\rtf'), 'Missing RTF signature'
        else:
            assert data.startswith((bytes.fromhex('d0cf11e0a1b11ae1'), b'{\\rtf')), 'Unknown DOC signature'
        extracted = run(['/usr/bin/textutil', '-convert', 'txt', '-stdout', '-encoding', 'UTF-8', path])
        assert extracted['returncode'] == 0 and extracted['stdout'], extracted
        return extracted['stdout'], details
    text = data.decode('utf8')
    if fmt == 'html':
        assert '<html' in text.lower() or '<!doctype' in text.lower(), 'Missing HTML document'
        parser = TextHTML()
        parser.feed(text)
        return ''.join(parser.parts), details
    return text, details


def content_checks(text, paragraphs):
    actual = normalize(text)
    missing = [p for p in paragraphs if normalize(p) not in actual]
    positions = [actual.find(normalize(p)) for p in paragraphs]
    exact = re.sub(r'\s+', '', text)
    return {'nonempty': bool(actual), 'expected_paragraph_count': len(paragraphs),
            'comparison_normalization': 'Unicode NFKC and whitespace removal',
            'exact_codepoint_paragraphs_present': all(re.sub(r'\s+', '', p) in exact for p in paragraphs),
            'missing_paragraph_count': len(missing), 'missing_paragraphs': missing,
            'paragraphs_in_order': positions == sorted(positions) and all(p >= 0 for p in positions),
            'chinese_preserved': all(t in actual for t in ['中文验收', '倒数第二段', '结束标记']),
            'first_and_last_markers_present': all(t in actual for t in ['BEGINFILEORBIT7429', 'ENDFILEORBIT9876'])}


def render_pdf(path, target_dir):
    target_dir.mkdir(parents=True, exist_ok=True)
    command = run(['pdftoppm', '-scale-to', '1000', '-png', path, target_dir / 'page'], timeout=90)
    assert command['returncode'] == 0, command
    images = sorted(target_dir.glob('page-*.png'))
    assert images, 'No rendered PNG pages'
    pixels = []
    for image in images:
        with Image.open(image) as im:
            grey = im.convert('L')
            histogram = grey.histogram()
            ink = sum(histogram[:230])
            pixels.append({'path': str(image), 'width': im.width, 'height': im.height,
                           'dark_pixel_count': ink, 'nonblank': ink > 100})
    assert all(p['nonblank'] for p in pixels), 'Blank rendered PDF page'
    return pixels


def make_fixtures(folder):
    folder.mkdir(parents=True, exist_ok=True)
    paragraphs = ['BEGIN FILEORBIT 7429 中文验收',
                  'English paragraph 0123456789; punctuation & < > and 中文数字 260315。',
                  'BOLD TEXT 13579 加粗文字保留测试',
                  'LIST ITEM ALPHA 中文列表第一项',
                  'LIST ITEM BETA 中文列表第二项']
    for i in range(1, 71):
        paragraphs.append(f'PARAGRAPH {i:03d} 中文验收段落。每段包含不同编号 {i * 113:05d}，用于检测截断、乱序和乱码。The quick brown fox crosses the quiet river safely.')
    paragraphs.extend(['倒数第二段 SECOND LAST 24680', 'END FILEORBIT 9876 结束标记'])
    fixtures = {}
    txt = folder / 'source.txt'
    txt.write_text('\n\n'.join(paragraphs) + '\n', encoding='utf8')
    fixtures['txt'] = txt
    md = folder / 'source.md'
    lines = ['# ' + paragraphs[0], paragraphs[1], '**' + paragraphs[2] + '**',
             '- ' + paragraphs[3] + '\n- ' + paragraphs[4], *paragraphs[5:]]
    md.write_text('\n\n'.join(lines) + '\n', encoding='utf8')
    fixtures['md'] = md
    ht = folder / 'source.html'
    body = [f'<h1>{html.escape(paragraphs[0])}</h1>', f'<p>{html.escape(paragraphs[1])}</p>',
            f'<p><strong>{html.escape(paragraphs[2])}</strong></p>',
            '<ul>' + ''.join(f'<li>{html.escape(p)}</li>' for p in paragraphs[3:5]) + '</ul>']
    body.extend(f'<p>{html.escape(p)}</p>' for p in paragraphs[5:])
    ht.write_text('<!DOCTYPE html><html><head><meta charset="utf-8"><style>body {font-family: Arial; font-size: 12pt} h1 {font-size: 18pt}</style></head><body>' + '\n'.join(body) + '</body></html>', encoding='utf8')
    fixtures['html'] = ht
    dx = folder / 'source.docx'
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(12)
    doc.add_paragraph(paragraphs[0], style='Title')
    doc.add_paragraph(paragraphs[1])
    doc.add_paragraph().add_run(paragraphs[2]).bold = True
    for p in paragraphs[3:5]:
        doc.add_paragraph(p, style='List Bullet')
    for p in paragraphs[5:]:
        doc.add_paragraph(p)
    doc.save(dx)
    fixtures['docx'] = dx
    for fmt in ['rtf', 'odt', 'doc']:
        dest = folder / ('source.' + fmt)
        generated = run(['/usr/bin/textutil', '-convert', fmt, '-output', dest, dx])
        (folder / ('generator-' + fmt + '.json')).write_text(json.dumps(generated, ensure_ascii=False, indent=2))
        assert generated['returncode'] == 0 and dest.exists(), generated
        fixtures[fmt] = dest
    source_checks = {}
    for fmt, path in fixtures.items():
        extracted, details = text_of(path, fmt)
        check = content_checks(extracted, paragraphs)
        assert check['missing_paragraph_count'] == 0 and check['paragraphs_in_order'], (fmt, check)
        source_checks[fmt] = {'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                              'checks': check, 'details': details}
    (folder / 'expected.json').write_text(json.dumps({'paragraphs': paragraphs, 'sources': source_checks}, ensure_ascii=False, indent=2))
    return fixtures, paragraphs


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--exe', required=True, type=Path)
    parser.add_argument('--root', required=True, type=Path)
    parser.add_argument('--timeout', type=int, default=45)
    args = parser.parse_args()
    root = args.root.resolve()
    root.mkdir(parents=True, exist_ok=True)
    exe = args.exe.resolve()
    fixtures, paragraphs = make_fixtures(root / 'fixtures')
    summary = {'executable': str(exe), 'executable_sha256': hashlib.sha256(exe.read_bytes()).hexdigest(),
               'description': 'Actual shipping CLI conversions; source and all 77 paragraph contents verified independently; PDF pages rasterized.',
               'pass_scope': 'Valid nonempty output with complete ordered text, preserved sources and visible results. Exact layout fidelity is reported separately, not implied by a content pass.',
               'cases': []}
    for source_format, fixture in fixtures.items():
        for target_format in ['pdf', 'docx', 'rtf', 'txt', 'html', 'odt', 'md']:
            if source_format == target_format:
                continue
            case_dir = root / 'cases' / (source_format + '-to-' + target_format)
            case_dir.mkdir(parents=True, exist_ok=True)
            source = case_dir / fixture.name
            shutil.copy2(fixture, source)
            original = hashlib.sha256(source.read_bytes()).hexdigest()
            case = {'source_format': source_format, 'target_format': target_format, 'status': 'failed',
                    'input': str(source), 'outputs': [], 'checks': {}, 'error': None}
            start = time.monotonic()
            try:
                command = run([exe, 'convert', source, '--to', target_format, '--json'], timeout=args.timeout)
                (case_dir / 'cli.json').write_text(json.dumps(command, ensure_ascii=False, indent=2))
                case['cli_returncode'] = command['returncode']
                result = json.loads(command['stdout'])
                case['outputs'] = result.get('outputs', [])
                case['cli_result'] = result
                assert command['returncode'] == 0 and result['succeededInputs'] == 1, result
                assert len(case['outputs']) == 1, result
                dest = Path(case['outputs'][0])
                assert dest.is_file() and dest.stat().st_size > 0, 'Missing or empty output file'
                text, details = text_of(dest, target_format)
                (case_dir / 'extracted.txt').write_text(text, encoding='utf8')
                checks = content_checks(text, paragraphs)
                checks['real_format_parse'] = True
                checks['source_unchanged'] = hashlib.sha256(source.read_bytes()).hexdigest() == original
                checks['output_visible'] = not bool(dest.stat().st_flags & 0x8000)
                case['checks'] = checks
                case['output_details'] = details
                if source_format == 'docx' and target_format == 'md':
                    case['fidelity_checks'] = {
                        'title_style_became_heading': bool(re.search(r'^#+ .*BEGIN FILEORBIT', text, re.M)),
                        'list_style_became_bullets': bool(re.search(r'^[-*+] .*LIST ITEM ALPHA', text, re.M)),
                        'direct_bold_preserved': '**BOLD TEXT 13579' in text}
                if target_format == 'pdf':
                    case['renders'] = render_pdf(dest, case_dir / 'render')
                assert checks['missing_paragraph_count'] == 0, f"Lost {checks['missing_paragraph_count']} expected paragraphs"
                assert checks['paragraphs_in_order'], 'Paragraph order changed'
                assert checks['source_unchanged'] and checks['output_visible'], 'Source modified or output hidden'
                case['status'] = 'passed'
            except Exception as e:
                case['error'] = repr(e)
            case['elapsed_seconds'] = round(time.monotonic() - start, 3)
            summary['cases'].append(case)
            summary['counts'] = {status: sum(c['status'] == status for c in summary['cases']) for status in ['passed', 'failed']}
            summary['distinct_directions'] = len({(c['source_format'], c['target_format']) for c in summary['cases']})
            (root / 'results.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2))
            print(f"{source_format}->{target_format}: {case['status']} {case['error'] or ''}", flush=True)
    print(json.dumps(summary['counts']), flush=True)


if __name__ == '__main__':
    main()
