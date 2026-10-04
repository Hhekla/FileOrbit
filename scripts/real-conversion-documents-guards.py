#!/usr/bin/env python3
"""Retain real DOCX fixtures for simple lists, unsupported numbering and image headers.

Run --fixtures-only first if the final application is still building, then use
--exe with the same --root. This script retains all evidence and never deletes it.
"""
import argparse
import importlib.util
import json
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches
from lxml import etree as ET
from PIL import Image, ImageDraw

spec = importlib.util.spec_from_file_location('extended_documents', Path(__file__).with_name('real-conversion-documents-extended.py'))
ext = importlib.util.module_from_spec(spec)
spec.loader.exec_module(ext)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def make_fixtures(root):
    folder = root / 'fixtures'
    folder.mkdir(parents=True, exist_ok=True)
    manifest = folder / 'manifest.json'
    if manifest.exists(): return json.loads(manifest.read_text())
    markers = ['BEGIN-LISTS-7429 中文列表验收', 'NUMBER-FIRST-1357 第一项',
               'NUMBER-SECOND-2468 第二项', 'BULLET-ITEM-3579 项目符号', 'END-LISTS-9876 结束标记']
    doc = Document()
    doc.add_paragraph(markers[0])
    doc.add_paragraph(markers[1], 'List Number')
    doc.add_paragraph(markers[2], 'List Number')
    doc.add_paragraph(markers[3], 'List Bullet')
    doc.add_paragraph(markers[4])
    simple = folder / 'simple-lists.docx'
    doc.save(simple)
    fixtures = [{'path': str(simple), 'variant': 'simple-list', 'markers': markers, 'expect_rejection': False}]
    with zipfile.ZipFile(simple) as package:
        contents = {name: package.read(name) for name in package.namelist()}
    styles = ET.fromstring(contents['word/styles.xml'])
    style = next(node for node in styles.findall(W + 'style') if node.get(W + 'styleId') == 'ListNumber')
    num_id = style.find('.//' + W + 'numId').get(W + 'val')
    for variant in ['roman', 'custom-marker', 'multilevel', 'start-override', 'level-override']:
        numbering = ET.fromstring(contents['word/numbering.xml'])
        num = next(node for node in numbering.findall(W + 'num') if node.get(W + 'numId') == num_id)
        abstract_id = num.find(W + 'abstractNumId').get(W + 'val')
        abstract = next(node for node in numbering.findall(W + 'abstractNum') if node.get(W + 'abstractNumId') == abstract_id)
        level = abstract.find(W + 'lvl')
        if variant == 'roman': level.find(W + 'numFmt').set(W + 'val', 'lowerRoman')
        elif variant == 'custom-marker': level.find(W + 'lvlText').set(W + 'val', 'Chapter %1)')
        elif variant == 'multilevel':
            kind = abstract.find(W + 'multiLevelType')
            if kind is not None: kind.set(W + 'val', 'multilevel')
            child = ET.SubElement(abstract, W + 'lvl', {W + 'ilvl': '1'})
            for tag, value in [('start', '1'), ('numFmt', 'decimal'), ('lvlText', '%1.%2.')]:
                ET.SubElement(child, W + tag, {W + 'val': value})
        else:
            override = ET.SubElement(num, W + 'lvlOverride', {W + 'ilvl': '0'})
            if variant == 'start-override': ET.SubElement(override, W + 'startOverride', {W + 'val': '7'})
            else:
                child = ET.SubElement(override, W + 'lvl', {W + 'ilvl': '0'})
                for tag, value in [('start', '3'), ('numFmt', 'decimal'), ('lvlText', '%1.')]:
                    ET.SubElement(child, W + tag, {W + 'val': value})
        path = folder / (variant + '.docx')
        with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as package:
            for name, data in contents.items():
                package.writestr(name, ET.tostring(numbering, xml_declaration=True, encoding='UTF-8') if name == 'word/numbering.xml' else data)
        fixtures.append({'path': str(path), 'variant': variant, 'markers': markers, 'expect_rejection': True, 'expected_error': '复杂列表编号'})
    image_path = folder / 'header.png'
    image = Image.new('RGB', (240, 80), '#147bd1')
    drawing = ImageDraw.Draw(image); drawing.rectangle((12, 12, 226, 67), fill='#eaab32')
    image.save(image_path)
    doc.sections[0].header.paragraphs[0].add_run().add_picture(str(image_path), width=Inches(2))
    path = folder / 'image-only-header.docx'; doc.save(path)
    with zipfile.ZipFile(path) as package:
        header = ET.fromstring(package.read('word/header1.xml'))
        assert header.find('.//' + W + 'drawing') is not None
        assert not any(node.text and node.text.strip() for node in header.iter(W + 't'))
    fixtures.append({'path': str(path), 'variant': 'image-only-header', 'markers': markers,
                     'expect_rejection': True, 'expected_error': '页眉'})
    for fixture in fixtures:
        path = Path(fixture['path'])
        with zipfile.ZipFile(path) as package: assert package.testzip() is None
        independent = Document(path)
        assert [p.text for p in independent.paragraphs] == markers
        fixture['source_sha256'] = ext.digest(path)
        fixture['source_checks'] = {'valid_docx_package': True, 'independent_python_docx_full_body': True}
    manifest.write_text(json.dumps(fixtures, ensure_ascii=False, indent=2))
    return fixtures


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--exe', type=Path)
    parser.add_argument('--root', required=True, type=Path)
    parser.add_argument('--fixtures-only', action='store_true')
    args = parser.parse_args()
    root = args.root.resolve()
    if (root / 'results.json').exists(): raise SystemExit('Choose a fresh --root; all previous evidence is retained.')
    fixtures = make_fixtures(root)
    if args.fixtures_only:
        print(f'Retained {len(fixtures)} independently readable DOCX fixtures in {root / "fixtures"}')
        return
    if not args.exe: parser.error('--exe is required unless --fixtures-only is used')
    exe = args.exe.resolve()
    report = {'executable': str(exe), 'executable_sha256': ext.digest(exe), 'fixtures': fixtures,
              'scope': 'Targeted final regression; simple lists must preserve text and numbering, unsupported structures must refuse without output. Not additional main-matrix directions.', 'cases': []}
    for fixture in fixtures:
        for target in ['pdf', 'txt', 'md']:
            folder = root / 'cases' / (fixture['variant'] + '-to-' + target)
            folder.mkdir(parents=True)
            source = folder / Path(fixture['path']).name; shutil.copy2(fixture['path'], source)
            before = ext.digest(source)
            case = {'source_format': 'docx', 'target_format': target, 'fixture_variant': fixture['variant'],
                    'input': str(source), 'status': 'failed', 'checks': {}, 'error': None, 'outputs': []}
            try:
                command = ext.run([exe, 'convert', source, '--to', target, '--json'])
                (folder / 'cli.json').write_text(json.dumps(command, ensure_ascii=False, indent=2))
                result = json.loads(command['stdout'])
                case['cli_returncode'] = command['returncode']; case['cli_result'] = result; case['outputs'] = result.get('outputs', [])
                if fixture['expect_rejection']:
                    case['checks'] = {'explicit_rejection': command['returncode'] != 0 and bool(result.get('failures')),
                                      'expected_reason': fixture['expected_error'] in json.dumps(result, ensure_ascii=False),
                                      'no_output': not case['outputs'] and not list(folder.glob('*.' + target))}
                    assert all(case['checks'].values()), case['checks']
                    case['status'] = 'expected_rejection'
                else:
                    assert command['returncode'] == 0 and result['succeededInputs'] == 1 and len(case['outputs']) == 1, result
                    output = Path(case['outputs'][0]); text, details = ext.basic.text_of(output, target)
                    (folder / 'extracted.txt').write_text(text)
                    actual = ext.basic.normalize(text)
                    positions = [actual.find(ext.basic.normalize(marker)) for marker in fixture['markers']]
                    case['checks'] = {'real_format_parse': True, 'full_body_in_order': all(p >= 0 for p in positions) and positions == sorted(positions),
                                      'first_number': '1.NUMBER-FIRST-1357' in actual, 'second_number': '2.NUMBER-SECOND-2468' in actual,
                                      'bullet_present': ('-BULLET-ITEM-3579' if target == 'md' else '•BULLET-ITEM-3579') in actual,
                                      'output_visible': not bool(output.stat().st_flags & 0x8000)}
                    assert all(case['checks'].values()), case['checks']
                    case['output_details'] = details
                    if target == 'pdf': case['renders'] = ext.basic.render_pdf(output, folder / 'render')
                    case['status'] = 'passed'
            except Exception as error: case['error'] = str(error)
            case['checks']['source_unchanged'] = ext.digest(source) == before
            if not case['checks']['source_unchanged']: case['status'] = 'failed'; case['error'] = 'Source changed'
            report['cases'].append(case)
            report['counts'] = {s: sum(c['status'] == s for c in report['cases']) for s in ['passed', 'expected_rejection', 'failed']}
            (root / 'results.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
            print(f'{fixture["variant"]} -> {target}: {case["status"]} {case["error"] or ""}', flush=True)


if __name__ == '__main__': main()
