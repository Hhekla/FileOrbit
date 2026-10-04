#!/usr/bin/env python3
"""Retained real-file document conversions, aliases, rich OOXML, long Chinese and corrupt input.

Use bundled Python with python-docx/Pillow/pypdf. --exe selects the built app;
--root must be a fresh evidence directory. This never deletes fixture/output files.
"""
import argparse
import base64
import hashlib
import html
import importlib.util
import json
import plistlib
import re
import shutil
import subprocess
import time
import zipfile
from urllib.parse import urlparse, unquote
from io import BytesIO
from pathlib import Path
import xml.etree.ElementTree as ET

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw

spec = importlib.util.spec_from_file_location('basic_documents', Path(__file__).with_name('real-conversion-documents.py'))
basic = importlib.util.module_from_spec(spec)
spec.loader.exec_module(basic)
TARGETS = ['pdf', 'docx', 'rtf', 'txt', 'html', 'odt', 'md']
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def run(command, timeout=120):
    try:
        p = subprocess.run(list(map(str, command)), capture_output=True, timeout=timeout)
        return {'command': list(map(str, command)), 'returncode': p.returncode,
                'stdout': p.stdout.decode(errors='replace'), 'stderr': p.stderr.decode(errors='replace')}
    except subprocess.TimeoutExpired as error:
        return {'command': list(map(str, command)), 'returncode': -999, 'stdout': '', 'stderr': str(error)}


def digest(path):
    h = hashlib.sha256()
    if path.is_dir():
        for item in sorted(path.rglob('*')):
            if item.is_file():
                h.update(str(item.relative_to(path)).encode())
                h.update(item.read_bytes())
    else:
        h.update(path.read_bytes())
    return h.hexdigest()


def make_fixtures(root):
    root.mkdir(parents=True, exist_ok=True)
    manifest_path = root / 'manifest.json'
    if manifest_path.exists():
        return json.loads(manifest_path.read_text())
    image = root / 'diagram.png'
    drawing = Image.new('RGB', (320, 160), '#ffffff')
    pen = ImageDraw.Draw(drawing)
    pen.rectangle((10, 10, 150, 150), fill='#1767a4')
    pen.ellipse((170, 15, 310, 155), fill='#ef6d20')
    pen.text((35, 60), 'IMAGE-4827', fill='white')
    drawing.save(image)
    basic_markers = ['BEGIN-ALIAS-7429 中文完整性', 'ALIAS-BODY-13579 正文第一段', 'ALIAS-END-9876 中文结束']
    fixtures = []
    for ext, content, same in [
        ('text', '\n\n'.join(basic_markers), 'txt'),
        ('markdown', '# ' + basic_markers[0] + '\n\n- ' + basic_markers[1] + '\n\n' + basic_markers[2], 'md'),
        ('htm', '<html><meta charset="utf-8"><body>' + ''.join('<p>' + x + '</p>' for x in basic_markers) + '</body></html>', 'html')]:
        path = root / ('alias.' + ext)
        path.write_text(content)
        fixtures.append({'path': str(path), 'source_format': ext, 'same_format': same, 'variant': 'extension-alias', 'markers': basic_markers})
    web_html = '<html><head><meta charset="utf-8"></head><body><h1>' + basic_markers[0] + '</h1><p>' + basic_markers[1] + '</p><img src="file:///fileorbit-fixture/diagram.png" width="320" height="160"><p>IMAGE-CAPTION-4827 图示</p><p>' + basic_markers[2] + '</p></body></html>'
    archive = {'WebMainResource': {'WebResourceData': web_html.encode(), 'WebResourceFrameName': '',
                                  'WebResourceMIMEType': 'text/html', 'WebResourceTextEncodingName': 'UTF-8',
                                  'WebResourceURL': 'file:///fileorbit-fixture/index.html'},
               'WebSubresources': [{'WebResourceData': image.read_bytes(), 'WebResourceMIMEType': 'image/png',
                                    'WebResourceURL': 'file:///fileorbit-fixture/diagram.png'}]}
    path = root / 'embedded.webarchive'
    path.write_bytes(plistlib.dumps(archive, fmt=plistlib.FMT_BINARY))
    fixtures.append({'path': str(path), 'source_format': 'webarchive', 'same_format': None,
                     'variant': 'embedded-image', 'markers': basic_markers[:2] + ['IMAGE-CAPTION-4827 图示'] + basic_markers[2:], 'image': True})
    rtfd = root / 'embedded.rtfd'
    helper = root / 'make-rtfd.swift'
    helper.write_text('''import AppKit
let out = URL(fileURLWithPath: CommandLine.arguments[1])
let png = URL(fileURLWithPath: CommandLine.arguments[2])
let text = NSMutableAttributedString(string: "BEGIN-ALIAS-7429 中文完整性\\nALIAS-BODY-13579 正文第一段\\n", attributes: [.font: NSFont(name: "Helvetica", size: 12)!])
let attachment = NSTextAttachment()
attachment.fileWrapper = try FileWrapper(url: png, options: [])
attachment.fileWrapper?.preferredFilename = "diagram.png"
text.append(NSAttributedString(attachment: attachment))
text.append(NSAttributedString(string: "\\nIMAGE-CAPTION-4827 图示\\nALIAS-END-9876 中文结束", attributes: [.font: NSFont(name: "Helvetica", size: 12)!]))
let wrapper = try text.fileWrapper(from: NSRange(location: 0, length: text.length), documentAttributes: [.documentType: NSAttributedString.DocumentType.rtfd])
try wrapper.write(to: out, options: [], originalContentsURL: nil)
''')
    generated = run(['/usr/bin/swift', '-module-cache-path', root / 'swift-cache', helper, rtfd, image])
    (root / 'rtfd-generator.json').write_text(json.dumps(generated, ensure_ascii=False, indent=2))
    assert generated['returncode'] == 0, generated
    fixtures.append({'path': str(rtfd), 'source_format': 'rtfd', 'same_format': None,
                     'variant': 'attached-image', 'markers': basic_markers[:2] + ['IMAGE-CAPTION-4827 图示'] + basic_markers[2:], 'image': True})
    rich = root / 'rich.docx'
    doc = Document()
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Title'].font.size = Pt(26)
    doc.styles['Heading 1'].font.size = Pt(18)
    doc.sections[0].header.paragraphs[0].text = 'HEADER-2468 页眉文本'
    doc.sections[0].footer.paragraphs[0].text = 'FOOTER-8642 页脚文本'
    markers = ['TITLE-7429 复杂文档验收', 'HEADING-1357 结构与样式', 'BODY-1111 加粗斜体与超链接',
               'BULLET-2222 无序列表第一项', 'BULLET-3333 无序列表第二项', 'NUMBER-4444 编号列表第一项',
               'NUMBER-5555 编号列表第二项', 'CELL-A1 项目', 'CELL-B1 数量', 'CELL-A2 样品', 'CELL-B2 12345',
               'IMAGE-CAPTION-4827 图示', 'EQUATION-LABEL-6666 公式', 'FOOTNOTE-ANCHOR-7777 脚注引用', 'END-RICH-9876 结束标记']
    doc.add_paragraph(markers[0], 'Title')
    doc.add_paragraph(markers[1], 'Heading 1')
    p = doc.add_paragraph(); p.add_run(markers[2]).bold = True
    p.add_run(' ITALIC-9999').italic = True
    for marker in markers[3:5]: doc.add_paragraph(marker, 'List Bullet')
    for marker in markers[5:7]: doc.add_paragraph(marker, 'List Number')
    table = doc.add_table(rows=2, cols=2); table.style = 'Table Grid'
    for cell, value in zip([cell for row in table.rows for cell in row.cells], markers[7:11]): cell.text = value
    doc.add_picture(str(image), width=Inches(3.2))
    doc.add_paragraph(markers[11])
    p = doc.add_paragraph(markers[12])
    math = OxmlElement('m:oMath'); mr = OxmlElement('m:r'); mt = OxmlElement('m:t'); mt.text = 'x+y=42'; mr.append(mt); math.append(mr); p._p.append(math)
    p = doc.add_paragraph(markers[13]); ref = OxmlElement('w:footnoteReference'); ref.set(qn('w:id'), '1'); p.add_run()._r.append(ref)
    doc.add_paragraph(markers[14])
    staged = root / 'rich-before-footnote.docx'; doc.save(staged)
    with zipfile.ZipFile(staged) as source, zipfile.ZipFile(rich, 'w', compression=zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == '[Content_Types].xml':
                data = data.replace(b'</Types>', b'<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
            elif item.filename == 'word/_rels/document.xml.rels':
                data = data.replace(b'</Relationships>', b'<Relationship Id="rIdFootnotes" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')
            target.writestr(item, data)
        target.writestr('word/footnotes.xml', f'<w:footnotes xmlns:w="{W}"><w:footnote w:id="1"><w:p><w:r><w:t>NOTE-9753 脚注正文保留</w:t></w:r></w:p></w:footnote></w:footnotes>')
    fixtures.append({'path': str(rich), 'source_format': 'docx', 'same_format': 'docx', 'variant': 'rich-layout',
                     'markers': markers, 'extra_markers': ['HEADER-2468 页眉文本', 'FOOTER-8642 页脚文本', 'x+y=42', 'NOTE-9753 脚注正文保留'], 'image': True})
    styled = Document()
    styled.styles['Normal'].font.size = Pt(11)
    styled.styles['Title'].font.size = Pt(26)
    styled.styles['Heading 1'].font.size = Pt(18)
    from docx.enum.style import WD_STYLE_TYPE
    inherited = styled.styles.add_style('FileOrbitTitle', WD_STYLE_TYPE.PARAGRAPH)
    inherited.base_style = styled.styles['Title']
    styled.add_paragraph(markers[0], 'FileOrbitTitle')
    styled.add_paragraph(markers[1], 'Heading 1')
    styled.add_paragraph(markers[2])
    for value in markers[3:5]: styled.add_paragraph(value, 'List Bullet')
    for value in markers[5:7]: styled.add_paragraph(value, 'List Number')
    table = styled.add_table(rows=2, cols=2); table.style = 'Table Grid'
    for cell, value in zip([cell for row in table.rows for cell in row.cells], markers[7:11]): cell.text = value
    styled.add_paragraph(markers[-1])
    path = root / 'styled-table.docx'; styled.save(path)
    fixtures.append({'path': str(path), 'source_format': 'docx', 'same_format': 'docx', 'variant': 'styled-table',
                     'markers': markers[:11] + markers[-1:]})
    image_doc = Document()
    image_doc.add_paragraph('IMAGE-DOCX-BEGIN-7429 中文正文')
    image_doc.add_picture(str(image), width=Inches(3))
    image_doc.add_paragraph('IMAGE-DOCX-END-9876 图片之后正文')
    path = root / 'image-only.docx'; image_doc.save(path)
    fixtures.append({'path': str(path), 'source_format': 'docx', 'same_format': 'docx', 'variant': 'image-only',
                     'markers': ['IMAGE-DOCX-BEGIN-7429 中文正文', 'IMAGE-DOCX-END-9876 图片之后正文'], 'image': True})
    longdoc = Document()
    longmarkers = ['LONG-TITLE-7429 长中文完整性验收']
    longdoc.add_heading(longmarkers[0], 0)
    for index in range(1, 401):
        text = f'LONG-{index:04d} 第{index}段中文。山川湖海与远行者的记录，保留所有句子、数字和标点，核查长文转换中的顺序与完整性。编号{index * 113}。'
        longmarkers.append(text); longdoc.add_paragraph(text)
    longmarkers.append('LONG-END-9876 四百段后的终止标记'); longdoc.add_paragraph(longmarkers[-1])
    path = root / 'long-chinese.docx'; longdoc.save(path)
    fixtures.append({'path': str(path), 'source_format': 'docx', 'same_format': 'docx', 'variant': 'long-chinese', 'markers': longmarkers})
    for ext, data in [('docx', b'PK\x03\x04broken archive'), ('rtf', b'{\\rtf1\\bin1000 only 8'), ('webarchive', b'bplist00broken')]:
        path = root / ('corrupt.' + ext); path.write_bytes(data)
        fixtures.append({'path': str(path), 'source_format': ext, 'same_format': ext, 'variant': 'corrupt', 'markers': [], 'expect_rejection': True})
    for fixture in fixtures:
        fixture['source_sha256'] = digest(Path(fixture['path']))
        if not fixture.get('expect_rejection'):
            source_path = Path(fixture['path'])
            if fixture['source_format'] == 'rtfd':
                decoded = run(['/usr/bin/textutil', '-convert', 'txt', '-stdout', '-encoding', 'UTF-8', source_path])
                assert decoded['returncode'] == 0, decoded
                source_text = decoded['stdout']
            elif fixture['source_format'] == 'webarchive':
                parsed = plistlib.loads(source_path.read_bytes())
                reader = basic.TextHTML(); reader.feed(parsed['WebMainResource']['WebResourceData'].decode())
                source_text = ''.join(reader.parts)
                with Image.open(BytesIO(parsed['WebSubresources'][0]['WebResourceData'])) as embedded: embedded.verify()
            elif fixture['source_format'] == 'docx':
                source_text, _ = basic.text_of(source_path, 'docx')
            elif fixture['source_format'] == 'htm':
                reader = basic.TextHTML(); reader.feed(source_path.read_text()); source_text = ''.join(reader.parts)
            else:
                source_text = source_path.read_text()
            missing = [m for m in fixture['markers'] if basic.normalize(m) not in basic.normalize(source_text)]
            assert not missing, ('Fixture failed independent source validation', fixture['path'], missing)
            fixture['source_checks'] = {'independently_decoded': True, 'missing_body_markers': missing}
    manifest_path.write_text(json.dumps(fixtures, ensure_ascii=False, indent=2))
    return fixtures


def media_count(path, fmt):
    if fmt in ['docx', 'odt']:
        with zipfile.ZipFile(path) as archive:
            return sum(n.startswith('word/media/' if fmt == 'docx' else 'Pictures/') for n in archive.namelist())
    if fmt == 'pdf':
        from pypdf import PdfReader
        return sum(len(p.images) for p in PdfReader(path).pages)
    if fmt == 'rtf': return path.read_bytes().count(b'\\pict')
    if fmt == 'html':
        images = re.findall(r'<img\b[^>]*\bsrc="([^"]*)"', path.read_text(), re.I)
        valid = 0
        for src in images:
            try:
                if src.startswith('data:image/'):
                    data = base64.b64decode(src.split(',', 1)[1])
                else:
                    parsed = urlparse(src)
                    candidate = Path(unquote(parsed.path)) if parsed.scheme == 'file' else path.parent / unquote(parsed.path)
                    if not candidate.resolve().is_relative_to(path.parent.resolve()): continue
                    data = candidate.read_bytes()
                with Image.open(BytesIO(data)) as image: image.verify()
                valid += 1
            except Exception: pass
        return valid
    return 0


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--exe', type=Path, required=True)
    parser.add_argument('--root', type=Path, required=True)
    parser.add_argument('--fixtures-only', action='store_true')
    parser.add_argument('--variants', help='comma-separated fixture variants to execute')
    args = parser.parse_args()
    root = args.root.resolve(); root.mkdir(parents=True, exist_ok=True)
    if (root / 'results.json').exists(): raise SystemExit('Choose a fresh --root; earlier evidence is retained.')
    fixtures = make_fixtures(root / 'fixtures')
    if args.fixtures_only: return
    exe = args.exe.resolve()
    report = {'executable': str(exe), 'executable_sha256': digest(exe), 'fixtures': fixtures,
              'pass_scope': 'Complete ordered fixture text, valid visible output, selected title/list/image checks. Plain-text table flattening and image omission are limited, not layout-fidelity passes.', 'cases': []}
    def save():
        report['counts'] = {s: sum(c['status'] == s for c in report['cases']) for s in ['passed', 'limited', 'failed', 'rejected', 'unsupported']}
        (root / 'results.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
    for fixture in fixtures:
        if args.variants and fixture['variant'] not in args.variants.split(','): continue
        for target in TARGETS:
            if target == fixture['same_format']: continue
            case_dir = root / 'cases' / (fixture['source_format'] + '-' + fixture['variant'] + '-to-' + target)
            case_dir.mkdir(parents=True)
            source = case_dir / Path(fixture['path']).name
            if Path(fixture['path']).is_dir(): shutil.copytree(fixture['path'], source)
            else: shutil.copy2(fixture['path'], source)
            before = digest(source)
            case = {'source_format': fixture['source_format'], 'target_format': target, 'fixture_variant': fixture['variant'],
                    'input': str(source), 'outputs': [], 'status': 'failed', 'checks': {}, 'error': None}
            command = run([exe, 'convert', source, '--to', target, '--json'], timeout=180)
            (case_dir / 'cli.json').write_text(json.dumps(command, ensure_ascii=False, indent=2))
            case['cli_returncode'] = command['returncode']
            try:
                result = json.loads(command['stdout']); case['cli_result'] = result; case['outputs'] = result.get('outputs', [])
                if fixture.get('expect_rejection'):
                    assert command['returncode'] != 0 and not case['outputs'], 'Corrupt input was accepted'
                    case['status'] = 'rejected'; case['checks']['explicit_rejection'] = True
                else:
                    assert command['returncode'] == 0 and result['succeededInputs'] == 1 and len(case['outputs']) == 1, result
                    dest = Path(case['outputs'][0]); text, details = basic.text_of(dest, target)
                    (case_dir / 'extracted.txt').write_text(text)
                    checks = case['checks']; normalized = basic.normalize(text)
                    checks['missing_body_markers'] = [m for m in fixture['markers'] if basic.normalize(m) not in normalized]
                    checks['missing_extra_markers'] = [m for m in fixture.get('extra_markers', []) if basic.normalize(m) not in normalized]
                    positions = [normalized.find(basic.normalize(m)) for m in fixture['markers']]
                    checks['body_in_order'] = positions == sorted(positions) and all(p >= 0 for p in positions)
                    checks['output_visible'] = not bool(dest.stat().st_flags & 0x8000)
                    checks['real_format_parse'] = True; checks['details'] = details
                    checks['embedded_image_count'] = media_count(dest, target)
                    if fixture.get('image'):
                        checks['image_preserved'] = checks['embedded_image_count'] > 0
                    if fixture['variant'] in ['rich-layout', 'styled-table'] and target == 'md':
                        checks['title_heading_preserved'] = bool(re.search(r'^#+ .*TITLE-7429', text, re.M))
                        checks['bullets_preserved'] = bool(re.search(r'^[-*+] .*BULLET-2222', text, re.M))
                        checks['ordered_list_preserved'] = bool(re.search(r'^\d+\. .*NUMBER-4444', text, re.M))
                    if fixture['variant'] in ['rich-layout', 'styled-table']:
                        if target == 'html': checks['table_structure_preserved'] = bool(re.search(r'<table\b', dest.read_text(), re.I))
                        elif target == 'rtf': checks['table_structure_preserved'] = b'\\trowd' in dest.read_bytes()
                        elif target == 'odt':
                            with zipfile.ZipFile(dest) as package: checks['table_structure_preserved'] = b'<table:table' in package.read('content.xml')
                        elif target in ['md', 'txt']: checks['table_structure_preserved'] = False
                    if target == 'pdf': checks['rendered_pages'] = basic.render_pdf(dest, case_dir / 'render')
                    assert not checks['missing_body_markers'] and checks['body_in_order'], 'Body text missing or reordered'
                    assert checks['output_visible'], 'Output hidden'
                    losses = bool(checks['missing_extra_markers']) or (bool(fixture.get('image')) and not checks['image_preserved'])
                    losses |= any(checks.get(k) is False for k in ['title_heading_preserved', 'bullets_preserved', 'ordered_list_preserved', 'table_structure_preserved'])
                    case['status'] = 'limited' if losses else 'passed'
                    if losses: case['error'] = 'Content/layout feature losses are recorded in checks; not a fidelity pass.'
            except Exception as error:
                case['error'] = str(error)
                if not fixture.get('expect_rejection') and case.get('cli_result', {}).get('failures') and not case['outputs']:
                    message = str(case['cli_result']['failures'])
                    if ('尚不能保留' in message or '不能保留文档附件图片' in message) and '已停止' in message:
                        case['status'] = 'unsupported'
            case['checks']['source_unchanged'] = digest(source) == before
            if not case['checks']['source_unchanged']:
                case['status'] = 'failed'; case['error'] = 'Source changed'
            report['cases'].append(case); save()
            print(f"{fixture['source_format']} {fixture['variant']} -> {target}: {case['status']} {case['error'] or ''}", flush=True)
    print(json.dumps(report['counts']), flush=True)


if __name__ == '__main__': main()
